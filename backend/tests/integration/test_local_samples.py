"""Optional local-only verification; portable test runs do not need company files."""

from collections import Counter
from hashlib import sha256
from io import BytesIO
import json
from time import perf_counter

from openpyxl import load_workbook
import pytest

from app.core.config import REPOSITORY_ROOT
from tests.integration.test_ingestion import create_cycle, upload


def test_original_workbook_ingestion_preserves_all_source_evidence(client, request):
    if not request.config.getoption("--run-local-samples"):
        pytest.skip("Confidential source verification is explicitly opt-in")
    samples = REPOSITORY_ROOT / "data/sample_files"
    inputs = [
        (samples / "Zf_LE1_Transformation_Check_Objects_Error_LOG.XLSX", "transformation", {"Detailed_LOG": 361}),
        (samples / "Upload errors.xlsx", "post_upload", {"Objects": 205, "Errors": 1226}),
    ]
    reference = samples / "ZF_LE1_Upload_Errors.xlsx"
    paths = [path for path, _, _ in inputs] + [reference]
    if not all(path.is_file() for path in paths):
        pytest.fail("Local company sample files are unavailable; omit --run-local-samples on a portable clone", pytrace=False)
    hashes = {path: sha256(path.read_bytes()).hexdigest() for path in paths}
    cycle_id, _ = create_cycle(client)
    collected_counts = {}
    started = perf_counter()
    for path, role, expected in inputs:
        content = path.read_bytes()
        response = upload(client, cycle_id, role, content, path.name)
        if response.status_code != 201:
            pytest.fail("Local ingestion failed; source content withheld from test output", pytrace=False)
        actual = response.json()["workbook"]["row_counts"]
        assert actual == expected
        collected_counts.update(actual)
    sheets = [
        (inputs[0][0], "Detailed_LOG", "transformation-findings", "Object Type"),
        (inputs[1][0], "Objects", "upload-objects", "Object Typ"),
        (inputs[1][0], "Errors", "upload-errors", "Object Typ"),
    ]
    for path, sheet_name, endpoint, type_header in sheets:
        records = []
        for offset in range(0, collected_counts[sheet_name], 200):
            records.extend(client.get(f"/api/v1/cycles/{cycle_id}/{endpoint}?limit=200&offset={offset}").json()["items"])
        if [record["excel_row_number"] for record in records] != list(range(2, collected_counts[sheet_name] + 2)):
            pytest.fail("Excel row provenance did not match the inspected source range", pytrace=False)
        with BytesIO(path.read_bytes()) as stream:
            book = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
            try:
                rows = list(book[sheet_name].iter_rows(values_only=True))
                headers = list(rows[0])
                for record, source in zip(records, rows[1:]):
                    if record["raw_payload"] != dict(zip(headers, source)):
                        pytest.fail("Raw evidence mismatch; source content withheld", pytrace=False)
                    if sheet_name == "Objects":
                        lookup = {str(header).strip(): value for header, value in zip(headers, source)}
                        if not (record["manual_status_raw"] == lookup["Manually added status"] and record["manual_comments_raw"] == lookup["Manually added comments"] and record["tool_status_raw"] == lookup["Status"]):
                            pytest.fail("Manual annotations and tool fields were not preserved independently", pytrace=False)
            finally:
                book.close()
        names = Counter(record["object_name_normalized"] for record in records if record["object_name_normalized"])
        expected_names = {"Detailed_LOG": 309, "Objects": 203, "Errors": 378}
        assert len(names) == expected_names[sheet_name]
        assert any(count > 1 for count in names.values())
    for path in paths:
        if sha256(path.read_bytes()).hexdigest() != hashes[path]:
            pytest.fail("An original workbook changed during verification", pytrace=False)
    print(json.dumps({"local_ingestion_counts": collected_counts, "elapsed_seconds": round(perf_counter() - started, 3), "original_hashes_unchanged": True}))
