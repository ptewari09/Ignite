from io import BytesIO

from openpyxl import load_workbook

from app.scripts.generate_demo_workbooks import post_upload_bytes, transformation_bytes, workbook_bytes


def mutate(content, change):
    with BytesIO(content) as stream:
        book = load_workbook(stream)
        change(book)
        return workbook_bytes(book)


def transformation_layout_bytes(sheet_name="Detailed_LOG"):
    def edit(book):
        sheet = book["Detailed_LOG"]
        # openpyxl treats a case-only rename as a name collision; use a
        # temporary name so the fixture has the exact intended Excel title.
        sheet.title = "Synthetic temporary title"
        sheet.title = sheet_name
    return mutate(transformation_bytes(), edit)


def raw_post_upload_bytes():
    """Fictional three-column object list and globally numbered technical log."""
    def edit(book):
        objects, errors = book["Objects"], book["Errors"]
        objects.title, errors.title = "Sheet1", "Sheet2"
        objects.delete_cols(4, 4)
        objects["C1"] = "Object Name".ljust(30)
        errors["D1"] = "Object Name".ljust(30)
        for row in (2, 3):
            objects.cell(row, 3).value = "z_demo_order".ljust(30)
        for row, sequence in enumerate((1, 2, 7, 11), 2):
            errors.cell(row, 2).value = sequence
    return mutate(post_upload_bytes(), edit)
