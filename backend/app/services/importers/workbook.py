from io import BytesIO
import warnings
from zipfile import ZipFile

from openpyxl import load_workbook

from app.core.config import Settings
from app.core.errors import AppError
from app.services.importers.transformation import parse_transformation
from app.services.importers.types import ImportWarning, ParsedWorkbook
from app.services.importers.upload_errors import parse_upload_errors
from app.services.importers.upload_objects import parse_upload_objects


def parse_workbook(content: bytes, role: str, settings: Settings) -> ParsedWorkbook:
    """Read bytes locally, without evaluating formulas or writing to a database."""
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > 10_000 or sum(m.file_size for m in members) > settings.max_workbook_uncompressed_bytes:
                raise AppError(413, "WORKBOOK_TOO_LARGE", "Expanded workbook exceeds the configured processing limit.")
            if any(m.flag_bits & 1 or m.filename.lower().endswith("vbaproject.bin") for m in members):
                raise AppError(422, "INVALID_WORKBOOK", "Encrypted or macro-enabled workbooks are not supported; supply an ordinary XLSX export.")
        with BytesIO(content) as source, warnings.catch_warnings(record=True) as captured:
            warnings.simplefilter("always", UserWarning)
            book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
            try:
                required = ("Detailed_LOG",) if role == "transformation" else ("Objects", "Errors")
                missing = [name for name in required if name not in book.sheetnames]
                if missing:
                    raise AppError(422, "MISSING_WORKSHEETS", f"Required worksheets are missing: {', '.join(missing)}.")
                result = ParsedWorkbook(sheets_found=book.sheetnames)
                adapters = [("Detailed_LOG", parse_transformation, "transformation")] if role == "transformation" else [
                    ("Objects", parse_upload_objects, "upload_objects"),
                    ("Errors", parse_upload_errors, "upload_errors"),
                ]
                for name, adapter, target in adapters:
                    parsed = adapter(book[name], settings)
                    setattr(result, target, parsed.records)
                    result.row_counts[name] = len(parsed.records)
                    result.object_counts[name] = parsed.object_count
                    result.warnings.extend(parsed.warnings)
                if captured:
                    result.warnings.append(ImportWarning(code="READER_WARNING", message="Workbook reader reported metadata warnings; parsed source rows were retained.", count=len(captured)))
                return result
            finally:
                book.close()
    except AppError:
        raise
    except Exception:
        # Library exceptions may contain source values or ZIP member paths.
        raise AppError(422, "INVALID_WORKBOOK", "The file is not a readable XLSX workbook. Export it again as .xlsx and retry.") from None
