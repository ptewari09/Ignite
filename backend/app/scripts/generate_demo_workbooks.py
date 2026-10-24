"""Generate fictional XLSX inputs locally; never read company sample data."""

import argparse
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill


def workbook_bytes(book: Workbook) -> bytes:
    with BytesIO() as output:
        book.save(output)
        book.close()
        return output.getvalue()


def transformation_bytes() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Detailed_LOG"
    sheet.append(["Status", "Object Type", "Object Name", "Error Category", "Message", "Suggested Action", "File Name"])
    sheet.append(["DEMO", " prog  ", " z_demo_order ", "OBJ_LOCKED", "Fictional object lock message.  ", "Review fictional evidence.", None])
    sheet.append([None] * 7)
    sheet.append(["DEMO", "PROG", "Z_DEMO_ORDER", "NO_TLAYER", "Fictional transport-layer message.", None, None])
    sheet.append([None, "CLAS", None, None, "Unidentified synthetic object.", None, None])
    sheet.cell(20, 1).fill = PatternFill(fill_type="solid", fgColor="EEEEEE")
    return workbook_bytes(book)


def post_upload_bytes() -> bytes:
    book = Workbook()
    objects = book.active
    objects.title = "Objects"
    objects.append(["Status Inf", "Object Typ", "Object Name", "Status    ", "Re Upload", "Manually added status", "Manually added comments"])
    objects.append(["DEMO ", "PROG  ", "Z_DEMO_ORDER", "UPLD_ERROR", "DEMO ", "Uploaded", "Synthetic review comment A."])
    objects.append([None, "PROG", "Z_DEMO_ORDER", "UPLD_ERROR", None, "Needs review", "Synthetic review comment B."])
    objects.append([None, "CLAS", "/DEMO/CL_TEST", "UPLD_ERROR", None, None, None])
    errors = book.create_sheet("Errors")
    errors.append(["Info  ", "S.No", "Object Typ", "Object Name", "Type", "Message        "])
    errors.append(["DEMO", 1, "PROG   ", "Z_DEMO_ORDER", "W   ", "Fictional request DEMO_TR_A is locked.   "])
    errors.append(["DEMO", 2, "    ", None, "I   ", "Fictional component context.   "])
    errors.append(["DEMO", 3, "PROG", "Z_DEMO_ORDER", "   ", "Fictional upload summary.   "])
    errors.append(["DEMO", 1, "CLAS", "/DEMO/CL_TEST", "E", "Fictional class upload message."])
    return workbook_bytes(book)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("data/demo"))
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    for filename, builder in [
        ("synthetic-transformation.xlsx", transformation_bytes),
        ("synthetic-post-upload.xlsx", post_upload_bytes),
    ]:
        target = args.output / filename
        if target.exists():
            print(f"Already exists; unchanged: {target}")
            continue
        with target.open("xb") as output:
            output.write(builder())
        print(f"Created fictional demo workbook: {target}")


if __name__ == "__main__":
    main()
